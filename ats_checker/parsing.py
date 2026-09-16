"""Text extraction + parseability risk checks.

Real ATS platforms parse a resume into structured fields (name, titles,
employers, dates, education, skills) via OCR + NLP *before* any keyword
matching happens. Parsing is binary in effect — a field either populates or
it doesn't. Two-column layouts, tables, images-as-text, and non-standard
section headers are the #1 reason qualified candidates get filtered out
before a human ever sees the resume. This module tries to catch those risks.
"""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

STANDARD_SECTIONS = {
    "experience": [
        r"\bwork experience\b", r"\bprofessional experience\b",
        r"\bemployment history\b", r"\bexperience\b",
    ],
    "education": [r"\beducation\b", r"\bacademic background\b"],
    "skills": [r"\bskills\b", r"\btechnical skills\b", r"\bcore competencies\b"],
    "summary": [r"\bsummary\b", r"\bobjective\b", r"\bprofile\b"],
}

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
# Phone formats vary too much for one strict pattern (+91 98765 43210,
# (555) 123-4567, 555-123-4567, +1.555.123.4567, ...). Instead, find runs of
# digits/spaces/dashes/dots/parens and check the digit count looks phone-like.
PHONE_CANDIDATE_RE = re.compile(r"[\d()+][\d()+\s.-]{7,}\d")

# Profile links a modern resume carries and ATS parsers extract into fields
# (LinkedIn, GitHub, portfolio...). Informational: their absence isn't a
# parse failure, but recruiters do look for them.
LINK_RE = re.compile(
    r"(?:https?://|www\.)[^\s,;|]+"
    r"|\blinkedin\.com/[^\s,;|]+"
    r"|\bgithub\.com/[^\s,;|]+"
    r"|\bgitlab\.com/[^\s,;|]+"
    r"|\bbehance\.net/[^\s,;|]+"
    r"|\bdribbble\.com/[^\s,;|]+"
    r"|\bmedium\.com/@?[^\s,;|]+",
    re.I,
)

# Date formats a resume might mix — consistency is a parseability signal.
DATE_MON_YEAR_RE = re.compile(r"\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+(?:19|20)\d{2}\b", re.I)
DATE_SLASH_RE = re.compile(r"\b(?:0?[1-9]|1[0-2])/(?:19|20)\d{2}\b")
DATE_ISO_RE = re.compile(r"\b(?:19|20)\d{2}-(?:0?[1-9]|1[0-2])\b")


def _has_phone_number(text: str) -> bool:
    for match in PHONE_CANDIDATE_RE.findall(text):
        digit_count = sum(c.isdigit() for c in match)
        if 9 <= digit_count <= 15:
            return True
    return False


@dataclass
class ParseResult:
    text: str
    file_type: str
    word_count: int = 0
    sections_found: list[str] = field(default_factory=list)
    sections_missing: list[str] = field(default_factory=list)
    has_email: bool = False
    has_phone: bool = False
    has_links: bool = False
    multi_column_risk: bool = False
    has_tables: bool = False
    warnings: list[str] = field(default_factory=list)

    def formatting_score(self) -> int:
        """0-100 heuristic score for how cleanly an ATS could parse this file."""
        score = 100
        score -= 15 * len(self.sections_missing)
        if not self.has_email:
            score -= 15
        if not self.has_phone:
            score -= 10
        if self.multi_column_risk:
            score -= 20
        if self.has_tables:
            score -= 10
        if self.word_count < 100:
            score -= 15  # likely an extraction failure, not a short resume
        return max(0, min(100, score))


def extract_text(path: str) -> tuple[str, str]:
    """Return (raw_text, file_type) for a resume file (.pdf, .docx, .txt)."""
    p = Path(path)
    suffix = p.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(p)
    if suffix == ".docx":
        return _extract_docx(p)
    if suffix in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="ignore"), "txt"
    raise ValueError(f"Unsupported file type: {suffix}. Use .pdf, .docx, or .txt")


def _extract_pdf(p: Path) -> tuple[str, str]:
    import pdfplumber

    text_parts = []
    multi_column_risk = False
    with pdfplumber.open(p) as pdf:
        for page in pdf.pages:
            text_parts.append(page.extract_text() or "")
            if _looks_multi_column(page):
                multi_column_risk = True
    text = "\n".join(text_parts)
    return text, ("pdf_multicol" if multi_column_risk else "pdf")


def _looks_multi_column(page) -> bool:
    """Cluster word left-edge (x0) positions; two dense, well-separated
    clusters spanning many lines is a strong multi-column signal — ATS
    parsers typically read left-to-right across the *whole* page width and
    will interleave the two columns into nonsense.
    """
    try:
        words = page.extract_words()
    except Exception:
        return False
    if len(words) < 40:
        return False

    x0s = sorted(w["x0"] for w in words)
    page_width = page.width or 1
    mid = page_width / 2

    left = [x for x in x0s if x < mid - page_width * 0.08]
    right = [x for x in x0s if x > mid + page_width * 0.08]

    # Multi-column: a real second cluster of word-starts well past the
    # midpoint, each cluster substantial (not just an indented bullet or two)
    return len(left) > len(x0s) * 0.25 and len(right) > len(x0s) * 0.25


def _extract_docx(p: Path) -> tuple[str, str]:
    import docx

    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    has_tables = len(d.tables) > 0
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = "\n".join(parts)
    return text, ("docx_tables" if has_tables else "docx")


def analyze(path: str | None = None, text: str | None = None) -> ParseResult:
    """Run parseability checks on a resume, given either a file path or raw text."""
    if path:
        raw_text, file_type = extract_text(path)
    elif text is not None:
        raw_text, file_type = text, "text"
    else:
        raise ValueError("Provide either path or text")

    result = ParseResult(text=raw_text, file_type=file_type)
    result.word_count = len(raw_text.split())
    result.has_email = bool(EMAIL_RE.search(raw_text))
    result.has_phone = _has_phone_number(raw_text)
    result.has_links = bool(LINK_RE.search(raw_text))
    result.has_tables = file_type == "docx_tables"
    result.multi_column_risk = file_type == "pdf_multicol"

    lower = raw_text.lower()
    for section, patterns in STANDARD_SECTIONS.items():
        if any(re.search(pat, lower) for pat in patterns):
            result.sections_found.append(section)
        else:
            result.sections_missing.append(section)

    if result.word_count < 100:
        result.warnings.append(
            "Very low word count extracted — the file may have failed to parse "
            "cleanly (scanned image, unusual encoding, or heavy graphics)."
        )
    if result.multi_column_risk:
        result.warnings.append(
            "Multi-column layout detected — many ATS parsers read straight "
            "across the page and will scramble two-column text."
        )
    if result.has_tables:
        result.warnings.append(
            "Tables detected in the .docx — some ATS parsers drop or misread "
            "table content entirely."
        )
    if not result.has_email:
        result.warnings.append("No email address detected — contact extraction may fail.")
    if not result.has_phone:
        result.warnings.append("No phone number detected — contact extraction may fail.")
    if not result.has_links:
        result.warnings.append(
            "No LinkedIn/GitHub/portfolio link detected — not a parse failure, but "
            "recruiters look for a LinkedIn profile and many ATS extract it into a field."
        )

    # Date-format consistency: mixing "Jan 2020" with "01/2020" styles across
    # roles is a real friction point for both parsers and recruiters.
    formats_used = []
    if DATE_MON_YEAR_RE.search(raw_text):
        formats_used.append("'Mon YYYY'")
    if DATE_SLASH_RE.search(raw_text):
        formats_used.append("'MM/YYYY'")
    if DATE_ISO_RE.search(raw_text):
        formats_used.append("'YYYY-MM'")
    if len(formats_used) > 1:
        result.warnings.append(
            "Mixed date formats detected (" + ", ".join(formats_used) + ") — "
            "use ONE consistent format for every role (e.g. 'Jan 2020 - Present')."
        )

    if result.sections_missing:
        result.warnings.append(
            "Missing standard section header(s): "
            + ", ".join(result.sections_missing)
            + " — use conventional headers (e.g. 'Experience', 'Education', 'Skills')."
        )

    return result


def word_frequencies(text: str) -> Counter:
    tokens = re.findall(r"[a-zA-Z][a-zA-Z0-9+.#/-]{1,}", text.lower())
    return Counter(tokens)
